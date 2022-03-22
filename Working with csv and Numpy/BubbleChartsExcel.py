# from openpyxl.chart import BubbleChart, Reference, Series
import openpyxl
from openpyxl.chart import Reference, BubbleChart, Series

# wb = openpyxl.Workbook()
# sheet = wb.active

# for i in range(10):
#     sheet.append([i * 2])

# values = Reference(sheet, min_col=1, min_row=1,
#                    max_col=1, max_row=10)

# chart = LineChart()
# chart.add_data(values)
# chart.title = "LINE CHART"
# chart.x_axis.title = "X_AXIS"
# chart.y_axis.title = "Y_AXIS"

# sheet.add_chart(chart, 'E2')
# wb.save('sample.xlsx')

wb = openpyxl.Workbook()

sheet = wb.active

rows = [
    ("Number of Products", "Sales in USD", "Market Share"),
    (14, 12200, 15),
    (20, 60000, 33),
    (18, 24400, 10),
    (22, 32000, 42)
]

for row in rows:
    sheet.append(row)

chart = BubbleChart()

# create data for plotting
xvalues = Reference(sheet, min_col=1,
                    min_row=2, max_row=5)

yvalues = Reference(sheet, min_col=2,
                    min_row=2, max_row=5)

size = Reference(sheet, min_col=3,
                 min_row=2, max_row=5)

# create a 1st series of data
series = Series(values=yvalues, xvalues=xvalues,
                zvalues=size, title="2013")

# add series data to the chart object
chart.series.append(series)

# set the title of the chart
chart.title = " BUBBLE-CHART "

# set the title of the x-axis
chart.x_axis.title = " X_AXIS "

# set the title of the y-axis
chart.y_axis.title = " Y_AXIS "

# add chart to the sheet
# the top-left corner of a chart
# is anchored to cell E2 .
sheet.add_chart(chart, "E2")

wb.save("BubbleChart.xlsx")
